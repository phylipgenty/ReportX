"""
Word (python-docx) Project Status Report in the FMDQ April 2025 format —
the editable companion to services/report_pdf.py, with the same sections,
tables, colours, legend and matrices. Executive Assistance Requests are
omitted (README §22).
"""
from __future__ import annotations
from io import BytesIO
from typing import List

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

from app.models.report import ProjectStatusReport
from app.services.report_pdf import LOGO, LOGO_SMALL, _us_date, _period_label, _milestone_date
from app.services.store import store

NAVY = "0B2D5B"
BORDER = "25375E"
RULE = "0B3080"
HEAD = "F2F2F2"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


# ---------------------------------------------------------------------------
# XML helpers
# ---------------------------------------------------------------------------
def _shade(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tc_pr.append(shd)


def _para_border(paragraph, side: str = "bottom", size: int = 24, color: str = RULE) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    borders.append(b)
    p_pr.append(borders)


def _field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {text} "
        run._r.append(el)


def _page_setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin, sec.bottom_margin = Inches(1.45), Inches(1)
    sec.header_distance = Inches(0.45)
    sec.different_first_page_header_footer = True
    sect_pr = sec._sectPr
    # navy page border on every page, as in the original
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "36")
        b.set(qn("w:space"), "24")
        b.set(qn("w:color"), BORDER)
        pg.append(b)
    sect_pr.append(pg)
    # cover is page 0, so the contents page is numbered 1 (matches the original)
    num = OxmlElement("w:pgNumType")
    num.set(qn("w:start"), "0")
    sect_pr.append(num)
    # ask Word to refresh the contents field when the file is opened
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    doc.settings.element.append(upd)


def render_docx(r: ProjectStatusReport) -> bytes:
    lookups = store.lookups()
    classification = str(store.setting("report_classification") or "")

    def opt(kind, value) -> dict:
        return next((o for o in lookups.get(kind, []) if o["value"] == str(value)), {})

    def label(kind, value) -> str:
        return opt(kind, value).get("label", str(value)) if value not in (None, "") else "N/A"

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    for name in ("Heading 1",):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Calibri", Pt(12), True
        st.font.color.rgb = RGBColor(0, 0, 0)
    _page_setup(doc)
    sec = doc.sections[0]

    # headers / footers: classification on every page, logo from page 2, page numbers after the cover
    for hdr in (sec.first_page_header, sec.header):
        p = hdr.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        run = p.add_run(classification)
        run.font.size = Pt(8)
        run.font.name = "Arial"
        if hdr is sec.header and LOGO_SMALL.exists():
            # the Header style has centre and right tab stops; two tabs reach the right edge
            p.add_run("\t\t").add_picture(str(LOGO_SMALL), width=Inches(2.2))
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(foot, "PAGE")

    def para(text="", bold=False, size=None, align=None, space_after=None):
        p = doc.add_paragraph()
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        if align is not None:
            p.alignment = align
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def cell_text(cell, text, bold=False, size=None, align=None):
        cell.text = ""
        lines = str(text).split("\n")
        p = cell.paragraphs[0]
        for i, line in enumerate(lines):
            run = p.add_run(line)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
            if i < len(lines) - 1:
                run.add_break()
        if align is not None:
            p.alignment = align
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    def cell_bullets(cell, items: List[str], size=9.5):
        items = [i for i in items if str(i).strip()] or ["N/A"]
        cell.text = ""
        for n, item in enumerate(items):
            p = cell.paragraphs[0] if n == 0 else cell.add_paragraph()
            run = p.add_run(f"▪  {item}")
            run.font.size = Pt(size)
            p.paragraph_format.space_after = Pt(6)

    def table(rows, widths, header=True, header_fill=HEAD):
        t = doc.add_table(rows=len(rows), cols=len(widths))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            for ci, value in enumerate(row):
                c = t.cell(ri, ci)
                c.width = Inches(6.5 * widths[ci])
                if value is not None:
                    cell_text(c, value, bold=header and ri == 0)
                if header and ri == 0 and header_fill:
                    _shade(c, header_fill)
        # rows never break across pages; the header row repeats on each page
        for ri, row in enumerate(t.rows):
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:cantSplit"))
            if header and ri == 0:
                tr_pr.append(OxmlElement("w:tblHeader"))
        return t

    def merge(a, b):
        """Merge cells without the blank paragraphs Word carries over from the merged cells."""
        cell = a.merge(b)
        for extra in cell.paragraphs[1:]:
            if not extra.text.strip():
                extra._p.getparent().remove(extra._p)
        return cell

    def keep_together(t):
        for row in t.rows[:-1]:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = True

    def font_size(t, size):
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(size)

    def lines_of(text: str) -> List[str]:
        return [l.strip(" •▪-") for l in (text or "").splitlines() if l.strip()]

    m = r.metadata
    placeholders = {o["value"]: o["label"] for o in lookups.get("date_placeholder", [])}

    # --- cover --------------------------------------------------------------
    for _ in range(5):
        para()
    if LOGO.exists():
        para(align=WD_ALIGN_PARAGRAPH.CENTER).add_run().add_picture(str(LOGO), width=Inches(5.8))
    for _ in range(5):
        para()
    _para_border(para(space_after=10))
    para("Project Status Report", bold=True, size=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para_border(para(m.project_title.upper(), bold=True, size=24, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10))
    for _ in range(3):
        para()
    _para_border(para(m.delivery_organisation or "", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER), side="top")
    month = para(_period_label(m.reporting_period.start_date, m.reporting_period.end_date),
                 bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para_border(month)
    month.add_run().add_break(WD_BREAK.PAGE)

    # --- contents (a Word TOC field, refreshed when the document opens) -------------
    para("Contents", size=13, space_after=14)
    _field(para(), 'TOC \\o "1-1" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # --- banner + metadata ------------------------------------------------------
    banner = doc.add_table(rows=1, cols=1)
    _shade(banner.cell(0, 0), NAVY)
    cell_text(banner.cell(0, 0), "Project Status Report", bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    banner.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = WHITE
    para()
    period = f"{_us_date(m.reporting_period.start_date)} – {_us_date(m.reporting_period.end_date)}"
    meta = table([
        ["Reporting period:", period, "Project title:", m.project_title],
        ["Date of report:", _us_date(m.date_of_report), "Delivery Manager:", m.delivery_manager or "N/A"],
        ["Report author:", m.report_author or "N/A", "Executive Sponsor:", m.executive_sponsor or "N/A"],
    ], (0.20, 0.27, 0.215, 0.315), header=False)
    for row in meta.rows:
        row.height = Inches(0.42)
    for ri in range(3):
        for ci in (0, 2):
            _shade(meta.cell(ri, ci), HEAD)
            meta.cell(ri, ci).paragraphs[0].runs[0].bold = True

    # --- 1. Executive summary -----------------------------------------------------
    doc.add_heading("1.  Executive Summary", level=1)
    ex = table([["Narrative Summary of Status", "Schedule:", "RAG Status", "Budget", "RAG Status", "Issues:", "RAG Status"],
                [None] * 7], (0.215, 0.19, 0.115, 0.105, 0.11, 0.155, 0.11), header_fill=None)
    cell_bullets(ex.cell(1, 0), lines_of(r.executive_summary) or lines_of(r.narrative))
    for i, key in enumerate(("schedule", "budget", "issues")):
        comment = getattr(r.rag, f"{key}_comment", "") or ""
        c_note, c_rag = ex.cell(1, 1 + i * 2), ex.cell(1, 2 + i * 2)
        if comment.strip():
            cell_bullets(c_note, lines_of(comment))
        else:
            cell_text(c_note, "N/A", align=WD_ALIGN_PARAGRAPH.CENTER)
        colour = opt("rag", getattr(r.rag, key)).get("report_color")
        if colour:
            _shade(c_rag, colour)
        else:
            cell_text(c_rag, "N/A", align=WD_ALIGN_PARAGRAPH.CENTER)

    # --- 2. Milestones ------------------------------------------------------------
    doc.add_heading("2.  Project Milestone Status Review", level=1)
    rows = [["S/N", "Project Milestones", "Status", "Baseline Completion Date", "Expected Completion Date", "Issues Exist (Yes/No)"]]
    for i, row in enumerate(r.milestone_status_review, start=1):
        expected = _milestone_date(row.expected_date, placeholders, "N/A")
        if row.actual_date and _us_date(row.actual_date):
            expected = f"{expected}\nActual: {_us_date(row.actual_date)}"
        rows.append([str(i), row.label, label("milestone_status", row.status),
                     _milestone_date(row.baseline_date, placeholders, "Yet to be determined"),
                     expected, "Yes" if row.has_issues else "No"])
    ms = table(rows, (0.115, 0.26, 0.13, 0.14, 0.13, 0.225))
    for ri in range(1, len(rows)):
        for ci in (0, 1):
            for run in ms.cell(ri, ci).paragraphs[0].runs:
                run.italic = True

    # --- 3. Planned activities ------------------------------------------------------
    doc.add_heading("3.  Status of Planned Activities", level=1)
    for title, items in (("Planned accomplishments in this period:", r.planned_activities.accomplishments),
                         ("Planned but not accomplished:", r.planned_activities.not_accomplished),
                         ("Planned actions for the next period:", r.planned_activities.next_period)):
        box = table([[title], [None]], (1.0,))
        cell_bullets(box.cell(1, 0), items, size=10.5)
        para()

    # --- 4. Issues + legend ---------------------------------------------------------
    doc.add_heading("4.  Project Issues Summary", level=1)
    rows = [["S/N", "Priority", "Issue Description", "Impact Summary (Milestone, Schedule Scope, Resources, Space…)", "Action Steps"]]
    for i, iss in enumerate(r.issues, start=1):
        impact = iss.impact_summary + (f"\n({', '.join(iss.impact_areas)})" if iss.impact_areas else "")
        rows.append([str(i), label("issue_priority", iss.priority), iss.description or "N/A", impact or "N/A", iss.action_steps or "N/A"])
    if not r.issues:
        rows.append(["1", "N/A", "N/A", "N/A", "N/A"])
    table(rows, (0.062, 0.107, 0.297, 0.268, 0.266))
    para()
    priorities = lookups.get("issue_priority", [])
    legend = table([["Legend", ""]] + [[o["label"] + (f" - {o['code']}" if o.get("code") else ""), o.get("description", "")]
                                      for o in priorities], (0.13, 0.87), header_fill="A6A6A6")
    merge(legend.cell(0, 0), legend.cell(0, 1))
    keep_together(legend)
    for i, o in enumerate(priorities, start=1):
        if o.get("report_color"):
            _shade(legend.cell(i, 0), o["report_color"])
        for run in legend.cell(i, 1).paragraphs[0].runs:
            run.italic = True
    font_size(legend, 8)

    # --- 5. Risks + matrices --------------------------------------------------------
    doc.add_heading("5.  Project Risk Summary", level=1)
    rows = [["S/N", "Priority", "Probability of Occurrence", "Risk Description",
             "Impact Summary (Milestone, Schedule Scope, Resources, Space…)", "Risk Score", "Response Strategy", "Status"]]
    for i, k in enumerate(r.risks, start=1):
        impact = str(k.impact) + (f"\n{k.impact_summary}" if k.impact_summary else "")
        rows.append([str(i), label("issue_priority", k.priority), str(k.probability), k.description or "N/A",
                     impact, str(k.risk_score), k.response_strategy or "N/A", label("risk_status", k.status)])
    if not r.risks:
        rows.append(["1"] + ["N/A"] * 7)
    font_size(table(rows, (0.055, 0.1, 0.125, 0.165, 0.135, 0.07, 0.185, 0.165)), 9)
    para()

    probs = lookups.get("probability", [])
    lm = table([["Likelihood Matrix", "", "", ""], ["Likelihood Score", "Descriptor", "Number of Occurrence", "Probability"]]
               + [[o["value"], o["label"], o.get("description", ""), o.get("range", "")] for o in probs],
               (0.095, 0.095, 0.66, 0.15), header_fill=NAVY)
    merge(lm.cell(0, 0), lm.cell(0, 3))
    keep_together(lm)
    for ci in range(4):
        _shade(lm.cell(1, ci), "B4C6E7" if ci < 2 else "BFBFBF")
        lm.cell(1, ci).paragraphs[0].runs[0].bold = True
    font_size(lm, 8)
    lm.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = WHITE
    lm.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    para()

    impacts = sorted(lookups.get("impact", []), key=lambda o: -int(o["value"]) if o["value"].isdigit() else 0)
    im = table([["Impact Matrix", "", ""], ["Impact Score", "Descriptor", "Project Impact"]]
               + [[o["value"], o["label"], o.get("description", "")] for o in impacts],
               (0.118, 0.103, 0.472), header_fill=NAVY)
    im.alignment = WD_TABLE_ALIGNMENT.LEFT
    merge(im.cell(0, 0), im.cell(0, 2))
    keep_together(im)
    for ci in range(3):
        _shade(im.cell(1, ci), "BFBFBF")
        im.cell(1, ci).paragraphs[0].runs[0].bold = True
    for i, o in enumerate(impacts, start=2):
        if o.get("report_color"):
            _shade(im.cell(i, 1), o["report_color"])
    font_size(im, 8)
    im.cell(0, 0).paragraphs[0].runs[0].font.color.rgb = WHITE
    im.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
